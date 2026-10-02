from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _set_run_font(run: Any, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)


def _clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _keep_with_next(paragraph: Any) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.widow_control = True


def _set_cell_margins(cell: Any, top: int = 60, start: int = 80, bottom: int = 60, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_entry(document: Document, entry: dict[str, Any]) -> None:
    heading = document.add_paragraph(style="CV Entry")
    heading.paragraph_format.tab_stops.add_tab_stop(Inches(7.15), WD_TAB_ALIGNMENT.RIGHT)
    title_run = heading.add_run(entry["title"])
    title_run.bold = True
    _set_run_font(title_run, "Arial", 10.3)
    if entry.get("date"):
        heading.add_run("\t")
        date_run = heading.add_run(entry["date"])
        _set_run_font(date_run, "Arial", 9.5)
    _keep_with_next(heading)

    if entry.get("subtitle"):
        subtitle = document.add_paragraph(style="CV Detail")
        run = subtitle.add_run(entry["subtitle"])
        run.italic = True
        _set_run_font(run, "Arial", 9.5)
        _keep_with_next(subtitle)

    for metadata in entry.get("metadata", []):
        paragraph = document.add_paragraph(style="CV Detail")
        run = paragraph.add_run(metadata)
        _set_run_font(run, "Arial", 9.5)

    for bullet in entry.get("bullets", []):
        paragraph = document.add_paragraph(style="CV Bullet")
        run = paragraph.add_run(bullet)
        _set_run_font(run, "Arial", 9.5)


def render_docx(context: dict[str, Any], template_path: Path, output_path: Path) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(
            f"DOCX template not found: {template_path}. Run scripts/build_docx_template.py first."
        )
    document = Document(str(template_path))
    _clear_body(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(context["name"])
    title_run.bold = True
    _set_run_font(title_run, "Arial", 18)

    if context.get("headline"):
        headline = document.add_paragraph(style="Subtitle")
        headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = headline.add_run(context["headline"])
        _set_run_font(run, "Arial", 10)

    if context.get("contact"):
        contact = document.add_paragraph(style="CV Contact")
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        values = [item["value"] for item in context["contact"]]
        run = contact.add_run(" | ".join(values))
        _set_run_font(run, "Arial", 9.2)

    for section in context["sections"]:
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.page_break_before = False
        run = heading.add_run(section["title"])
        run.bold = True
        _set_run_font(run, "Arial", 11.5)
        _keep_with_next(heading)

        if section["id"] == "research_interests":
            paragraph = document.add_paragraph(style="CV Detail")
            run = paragraph.add_run(", ".join(section["items"]))
            _set_run_font(run, "Arial", 9.7)
            continue

        for entry in section["items"]:
            _add_entry(document, entry)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = f"{context['name']} Academic CV"
    document.core_properties.subject = context["profile"]
    document.core_properties.author = context["name"]
    document.save(str(output_path))
    return output_path


def assert_docx_structure(path: Path) -> None:
    document = Document(str(path))
    nonempty = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    if not nonempty:
        raise ValueError(f"Generated DOCX contains no visible paragraphs: {path}")
    if len(document.sections) != 1:
        raise ValueError(f"Generated DOCX should have one document section: {path}")
    section = document.sections[0]
    if section.page_width != Inches(8.5) or section.page_height != Inches(11):
        raise ValueError(f"Generated DOCX is not US Letter portrait: {path}")

